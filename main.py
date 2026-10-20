from fastapi import FastAPI, File, UploadFile
from fastapi.middleware.cors import CORSMiddleware
import chardet
import hashlib
import json
import os
import re
import requests
import sqlite3
import tempfile
import time
import observability
from pypdf import PdfReader
from docx import Document as DocxDocument
from pydantic import BaseModel
from dotenv import load_dotenv
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_openai import ChatOpenAI
from langchain_core.tools import tool
from langchain.agents import create_agent
from langchain_chroma import Chroma             
from langchain_core.embeddings import Embeddings
from langchain_core.documents import Document
from langgraph.checkpoint.memory import InMemorySaver
from langgraph.checkpoint.sqlite import SqliteSaver
from pathlib import Path
from chromadb.utils.embedding_functions.onnx_mini_lm_l6_v2 import ONNXMiniLM_L6_V2

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ONNXMiniLM_L6_V2.DOWNLOAD_PATH = (
    Path(BASE_DIR) / "models" / "onnx_models" / ONNXMiniLM_L6_V2.MODEL_NAME
)

# ==================== 基础配置 ====================
load_dotenv()

# ==================== 大模型配置（模型层可插拔） ====================
# 换模型 / 换供应商只需改 .env，代码不用动。兼容旧的 ARK_* 变量名。
LLM_BASE_URL = os.getenv("LLM_BASE_URL", "https://ark.cn-beijing.volces.com/api/v3").rstrip("/")
LLM_API_KEY = os.getenv("LLM_API_KEY") or os.getenv("ARK_API_KEY", "")
LLM_MODEL = os.getenv("LLM_MODEL") or os.getenv("ARK_MODEL", "ep-20261008111305-gwhdt")
# 排序/判断类任务可以换更轻量的模型来降本提速；不填则与主模型一致
RERANK_MODEL = os.getenv("RERANK_MODEL") or LLM_MODEL
# 关闭思考链的请求参数：各厂商字段名不同（通义是 enable_thinking），留空表示不发送；
# 若厂商不接受该参数，call_ai 会自动去掉后重试一次。
_NO_THINK_RAW = os.getenv("LLM_DISABLE_THINKING_JSON", '{"thinking": {"type": "disabled"}}')
try:
    NO_THINK_PARAMS = json.loads(_NO_THINK_RAW) if _NO_THINK_RAW.strip() else {}
except ValueError:
    NO_THINK_PARAMS = {}
app = FastAPI(title="RAG 知识库问答助手")

# 跨域配置
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)



# ==================== 基础接口 ====================
@app.get("/")
def hello():
    return {"message": "RAG 知识库问答服务启动成功！"}

# 模块级缓存：ChatModel 只创建一次，复用底层 HTTP 连接池
_llm = None


def get_llm():
    """懒加载大模型客户端（LangChain ChatOpenAI，对接任意 OpenAI 兼容接口）"""
    global _llm
    if _llm is None:
        if not LLM_API_KEY:
            raise RuntimeError("未配置 LLM_API_KEY（或 ARK_API_KEY），请在项目根目录的 .env 中填写后重启后端")
        _llm = ChatOpenAI(
            model=LLM_MODEL,
            api_key=LLM_API_KEY,
            base_url=LLM_BASE_URL,
            temperature=0.1,
            timeout=120,
            max_retries=2,
        )
    return _llm


def call_ai(prompt: str, label: str = "chat", no_thinking: bool = False, model: str | None = None):
    """调用大模型（带耗时 / token / 成本埋点；支持任意 OpenAI 兼容接口）"""
    if not LLM_API_KEY:
        raise RuntimeError("未配置 LLM_API_KEY（或 ARK_API_KEY），请在项目根目录的 .env 中填写后重启后端")

    url = f"{LLM_BASE_URL}/chat/completions"

    headers = {
        "Authorization": f"Bearer {LLM_API_KEY}",
        "Content-Type": "application/json"
    }

    data = {
        "model": model or LLM_MODEL,
        "messages": [
            {"role": "system", "content": "你是一个助手，严格按照要求的格式输出。"},
            {"role": "user", "content": prompt}
        ],
        "temperature": 0.1
    }

    extra = dict(NO_THINK_PARAMS) if no_thinking else {}
    data.update(extra)

    t0 = time.perf_counter()

    def elapsed():
        return (time.perf_counter() - t0) * 1000

    def post():
        try:
            return requests.post(url, headers=headers, json=data, timeout=(10, 120)), None
        except requests.exceptions.Timeout:
            return None, "timeout"
        except requests.exceptions.RequestException as e:
            return None, str(e)

    response, err = post()
    if response is not None and response.status_code >= 400 and extra:
        for k in extra:
            data.pop(k, None)
        response, err = post()

    if err == "timeout":
        observability.record_llm_call(label, elapsed(), ok=False, error="timeout")
        raise RuntimeError("调用大模型超时，请稍后重试")
    if err:
        observability.record_llm_call(label, elapsed(), ok=False, error=err)
        raise RuntimeError(f"调用大模型失败：{err[:120]}")

    if response.status_code != 200:
        observability.record_llm_call(label, elapsed(), ok=False, error=f"HTTP {response.status_code}")
        raise RuntimeError(f"大模型接口返回 {response.status_code}：{response.text[:200]}")

    try:
        result = response.json()
        usage = result.get("usage") or {}
        observability.record_llm_call(label, elapsed(),
                                      usage.get("prompt_tokens"), usage.get("completion_tokens"))
        return result["choices"][0]["message"]["content"]
    except (ValueError, KeyError, IndexError, TypeError):
        observability.record_llm_call(label, elapsed(), ok=False, error="bad response")
        raise RuntimeError(f"大模型返回格式异常：{response.text[:200]}")


# ==================== RAG 知识库 ====================
def read_text_auto(file_path):
    """按探测到的编码读文本，避免 GBK 文档被 errors='ignore' 吃掉中文"""
    with open(file_path, 'rb') as f:
        raw = f.read()
    encoding = chardet.detect(raw[:10000]).get('encoding') or 'utf-8'
    try:
        return raw.decode(encoding, errors='ignore')
    except (LookupError, UnicodeDecodeError):
        return raw.decode('utf-8', errors='ignore')

CHROMA_PATH = "./chroma_db"
os.makedirs(CHROMA_PATH, exist_ok=True)


class ChromaONNXEmbeddings(Embeddings):
    """把 Chroma 自带的 ONNX 嵌入函数包装成 LangChain 的 Embeddings 接口。

    好处：直接复用已经下载到 models/ 的本地模型，不必再引入
    sentence-transformers + torch 那套重依赖。
    """

    def __init__(self):
        self._ef = ONNXMiniLM_L6_V2()

    def embed_documents(self, texts):
        return [[float(x) for x in v] for v in self._ef(list(texts))]

    def embed_query(self, text):
        return [float(x) for x in self._ef([text])[0]]


_kb_init_error = ""
if os.getenv("KB_DISABLED") == "1":
    # 显式关闭知识库：CI 环境没有本地嵌入模型与向量库（都被 gitignore），
    # 关掉后 import 更快、行为也更确定；依赖知识库的测试会自动跳过
    vectorstore = None
    _kb_init_error = "已通过环境变量 KB_DISABLED=1 禁用"
else:
    try:
        vectorstore = Chroma(
            collection_name="knowledge_base",
            embedding_function=ChromaONNXEmbeddings(),
            persist_directory=CHROMA_PATH,
        )
    except Exception as e:
        vectorstore = None
        _kb_init_error = str(e)

# 会话记忆：必须放在模块级，跨请求共享，否则每次请求都"失忆"
# 用 SQLite 持久化 —— 重启服务后对话记录依然在。
# check_same_thread=False 是必须的：FastAPI 的同步接口跑在线程池里，
# 不同请求可能落在不同线程，而 SQLite 连接默认只允许创建它的线程使用。
CHAT_DB_PATH = os.getenv("CHAT_DB_PATH") or os.path.join(BASE_DIR, "chat_memory.sqlite3")
try:
    _chat_conn = sqlite3.connect(CHAT_DB_PATH, check_same_thread=False)
    _rag_checkpointer = SqliteSaver(_chat_conn)
    _rag_checkpointer.setup()          # 建表（checkpoints / writes），幂等可重复调用
except Exception:
    # 持久化不可用时退回内存记忆，保证服务仍能启动
    _rag_checkpointer = InMemorySaver()


def kb_unavailable():
    return {"success": False, "error": f"知识库初始化失败：{_kb_init_error}"}


def split_text(text, chunk_size=500, chunk_overlap=50):
    """文本切片（LangChain RecursiveCharacterTextSplitter 实现）"""
    if not text:
        return []
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        length_function=len,
        separators=["\n\n", "\n", "。", "！", "？", " ", ""],
    )
    return splitter.split_text(text)


# ==================== RAG接口：文档上传入库 ====================
@app.post("/rag/upload")
def rag_upload(file: UploadFile = File(...)):
    if vectorstore is None:
        return kb_unavailable()
    try:
        suffix = os.path.splitext(file.filename)[1].lower()
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            tmp.write(file.file.read())
            tmp_path = tmp.name

        # 解析文档文本
        text = ""
        if suffix == ".pdf":
            reader = PdfReader(tmp_path)
            for page in reader.pages:
                text += page.extract_text() + "\n\n"
        elif suffix == ".docx":
            doc = DocxDocument(tmp_path)      
            for para in doc.paragraphs:
                text += para.text + "\n"
        elif suffix in [".txt", ".md", ".csv"]:
            text = read_text_auto(tmp_path)
        else:
            return {"success": False, "error": f"不支持的文件格式：{suffix}"}

        # 切片并存库
        chunks = split_text(text)
        if not chunks:
            return {"success": False, "error": "文档中没有提取到可用文本（扫描版 PDF 常见）"}

        digest = hashlib.md5(text.encode("utf-8", errors="ignore")).hexdigest()[:8]
        ids = [f"{file.filename}-{digest}-{i}" for i in range(len(chunks))]
        vectorstore.add_documents(
            documents=[
                Document(page_content=c, metadata={"source": file.filename})
                for c in chunks
            ],
            ids=ids,
        )

        return {
            "success": True,
            "filename": file.filename,
            "chunks": len(chunks)
        }

    except Exception as e:
        return {"success": False, "error": str(e)}
    finally:
        if 'tmp_path' in dir():
            try:
                os.unlink(tmp_path)
            except:
                pass


# ==================== RAG接口：列出知识库文档 ====================
@app.get("/rag/documents")
def rag_documents():
    """列出知识库里已有的文档（按来源聚合片段数）"""
    if vectorstore is None:
        return kb_unavailable()
    try:
        data = vectorstore.get()
        counts = {}
        for meta in data.get("metadatas") or []:
            src = (meta or {}).get("source")
            if src:
                counts[src] = counts.get(src, 0) + 1
        docs = [{"source": s, "chunks": n} for s, n in sorted(counts.items())]
        return {
            "success": True,
            "documents": docs,
            "total_chunks": sum(counts.values()),
        }
    except Exception as e:
        return {"success": False, "error": str(e)}

# 两阶段检索参数：先宽召回，再由大模型重排
RECALL_K = 10    # 向量召回数量（候选越少，重排推理量越小）
FINAL_K = 3     # 重排后真正喂给模型的片段数



# 两阶段检索参数：先宽召回，再由大模型重排
_RERANK_PROMPT = """判断每段候选资料与用户提问的相关性，按相关性从高到低输出所有片段的编号。

只输出编号，用英文逗号分隔，不要解释、不要换行。
例如：2,7,0,5

用户提问：{question}

候选资料：
{snippets}
"""


def _parse_ranked(raw, n):
    """取出候选编号；兼容 '3,7,1' 与 {"ranked":[...]} 两种格式"""
    text = (raw or "").strip()
    if not text:
        return []
    start, end = text.find("{"), text.rfind("}")
    if start != -1 and end != -1:
        try:
            ranked = json.loads(text[start:end + 1]).get("ranked")
            if isinstance(ranked, list):
                return [i for i in ranked if isinstance(i, int) and 0 <= i < n]
        except Exception:
            pass
    # 退化处理：把输出里的整数当编号。
    # 注意正则要带 -?，否则 "-3" 会被 \d+ 抓成 "3"，负号丢失 → 静默取到错误片段
    nums = [int(x) for x in re.findall(r"-?\d+", text)]
    return [i for i in nums if 0 <= i < n][:n]

def rerank_documents(question, docs, top_n=FINAL_K):
    """第二阶段：用大模型对召回结果做相关性重排。

    任何异常（模型输出不合法、接口超时）都回退到向量排序，
    保证重排只可能变好、不会让检索挂掉。
    """
    docs = list(docs)
    if len(docs) <= 1:
        return docs[:top_n]
    snippets = "\n\n".join(f"[{i}] {d.page_content[:150]}" for i, d in enumerate(docs))
    try:
        raw = call_ai(_RERANK_PROMPT.format(question=question, snippets=snippets, top_n=top_n),
                      label="rerank", no_thinking=True, model=RERANK_MODEL)
        order = _parse_ranked(raw, len(docs))
        if not order:
            # 解析失败：记为失败指标，避免"表面正常、实际退化成向量排序"
            observability.record_llm_call("rerank-fallback", 0, ok=False, error="解析失败，回退向量排序")
            return docs[:top_n]
        return [docs[i] for i in order[:top_n]]
    except Exception as e:
        observability.record_llm_call("rerank-fallback", 0, ok=False, error=str(e)[:120])
        return docs[:top_n]

def build_list_documents_tool():
    """知识库的"元数据"工具：列出真实文档清单。

    为什么单独做一个工具：向量检索只能捞回 top-N 片段，模型会拿这几个片段
    反推"知识库一共几份文档"，从而过度概括（实测它把 6 份说成 2 份）。
    元数据类问题应当由工具直接给准确答案，而不是让检索去猜。
    """

    @tool("list_knowledge_documents", parse_docstring=True)
    def list_knowledge_documents() -> str:
        """列出知识库里已有的全部文档及其片段数。

        当用户询问"知识库里有哪些文档 / 有什么资料 / 一共几份文档"时使用本工具。
        """
        if vectorstore is None:
            return "知识库不可用。"
        data = vectorstore.get(include=["metadatas"])
        counts = {}
        for meta in data.get("metadatas") or []:
            src = (meta or {}).get("source")
            if src:
                counts[src] = counts.get(src, 0) + 1
        if not counts:
            return "知识库还是空的，还没有上传任何文档。"
        lines = [f"- {s}（{n} 个片段）" for s, n in sorted(counts.items())]
        return f"知识库中共有 {len(counts)} 份文档：\n" + "\n".join(lines)

    return list_knowledge_documents



def build_retriever_tool(captured, source=None, use_rerank=True):
    """构建 RAG 检索工具（两阶段：向量宽召回 + LLM 重排）。

    captured：把检索到的原文回传给接口层，用于溯源展示
    source：限定只在某一份文档里检索，避免"问 A 文档却答成 B 文档"
    """
    recall_kwargs = {"k": RECALL_K}
    if source:
        recall_kwargs["filter"] = {"source": source}
    wide_retriever = vectorstore.as_retriever(search_kwargs=recall_kwargs)

    @tool("search_knowledge_base", parse_docstring=True)
    def search_knowledge_base(query: str) -> str:
        """在本地知识库中检索与问题最相关的资料片段。

        Args:
            query: 检索用的自然语言问题
        """
        candidates = wide_retriever.invoke(query)
        docs = rerank_documents(query, candidates, top_n=FINAL_K) if use_rerank else candidates[:FINAL_K]
        captured["docs"] = [d.page_content for d in docs]
        if not docs:
            return "知识库中没有检索到相关内容。"
        return "\n\n".join(d.page_content for d in docs)

    return search_knowledge_base


class RagChatRequest(BaseModel):
    question: str
    session_id: str = "default"
    source: str | None = None
    # 是否启用 LLM 重排：精度更高，但每次多一次模型调用
    use_rerank: bool = True


_rag_system_prompt = (
    "你是知识库问答助手。\n"
    "1. 知识库里已经上传了文档。用户问“这份文档/这份文件/这个项目讲了什么”时，"
    "“这份/这个”指的就是知识库里的文档（若系统限定了检索范围，就是那份文档），"
    "必须直接调用 search_knowledge_base 检索，然后基于检索结果回答；不要反问“是哪份文档”。\n"
    "2. 用户问“知识库里有哪些文档 / 有什么资料 / 一共几份文档”时，"
    "必须调用 list_knowledge_documents 拿真实清单，**不要用检索到的几个片段反推文档总数**。\n"
    "3. 只要问题涉及知识库内容，都先检索，再严格基于检索结果回答，不要编造。\n"
    "4. 检索到的资料必须与问题主题相关才可使用。如果检索结果与用户问的主题**无关**"
    "（例如问的是 A 主题、检索到的却是 B 主题的资料），直接回答“知识库中没有与 XXX 相关的内容”，"
    "**绝不能用无关资料拼凑答案**。\n"
    "5. 只有用户明确询问对话本身（如“我刚才问的是什么”）时才不检索，直接结合上下文回答。\n"
    "6. 回答用中文。\n"
)


@app.post("/rag/chat")
def rag_chat(req: RagChatRequest):
    if vectorstore is None:
        return kb_unavailable()
    try:
        # 知识库为空时先给友好提示，省掉两次模型调用
        if vectorstore._collection.count() == 0:
            return {"success": False, "error": "知识库还是空的，请先上传文档"}

        # 检索工具：两阶段检索（宽召回 + 重排），按需限定文档范围
        collected = {"docs": []}
        search_tool = build_retriever_tool(collected, source=req.source, use_rerank=req.use_rerank)

        agent = create_agent(
            model=get_llm(),
            tools=[search_tool, build_list_documents_tool()],
            checkpointer=_rag_checkpointer,
            system_prompt=_rag_system_prompt,
        )

        # LangGraph 1.2+ 的 invoke 方法不接受 callbacks，改为直接在 agent 的 with_config 中设置
        result = agent.with_config(
            callbacks=[observability.MetricsCallback("rag")],
            configurable={"thread_id": req.session_id}
        ).invoke(
            {"messages": [{"role": "user", "content": req.question}]}
        )

        return {
            "success": True,
            "answer": result["messages"][-1].content,
            "related_docs": collected["docs"],
        }
    except Exception as e:
        import traceback
        traceback.print_exc()
        return {"success": False, "error": str(e)}

        # ==================== 可观测性：指标查询 ====================

@app.get("/metrics")
def get_metrics():
    """大模型调用指标：次数 / token / 延迟（均值、P95）/ 估算成本"""
    return observability.snapshot()